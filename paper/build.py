"""Build the paper into the official Apart template.

  python paper/build.py            -> paper/out/undercover_agents.docx (+ .pdf if LibreOffice is on PATH
                                       or SOFFICE is set) and paper/out/number_sources.md

Every number is resolved from the committed analysis outputs by paper_numbers.render; the build stops
on a missing key. The template's own paragraphs are used as formatting prototypes.
"""
from __future__ import annotations

import copy
import os
import re
import subprocess
import sys
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import content  # noqa: E402
import paper_numbers as N  # noqa: E402

TEMPLATE = Path(os.environ.get("APART_TEMPLATE", HERE / "template" / "Apart_Research_hackathon_submission_template.docx"))
OUT = HERE / "out"
FIGS = HERE / "figures"
FONT = "Old Standard TT"
BODY_PT = 10.5
W = qn


def text_of(el) -> str:
    return "".join(t.text or "" for t in el.iter(W("w:t")))


def find_par(body, needle):
    for p in body.iter(W("w:p")):
        if text_of(p).strip().startswith(needle):
            return p
    raise KeyError(needle)


def strip_runs(p):
    for child in list(p):
        if child.tag in (W("w:r"), W("w:hyperlink"), W("w:bookmarkStart"), W("w:bookmarkEnd")):
            p.remove(child)
    return p


MINUS_FONT = "Times New Roman"  # Old Standard TT draws U+2212 as a long dash and lacks Greek kappa
FALLBACK_CHARS = "\u2212\u03ba"   # characters set in MINUS_FONT
URL_RE = re.compile(r"(https?://[^\s]+[^\s.,;)])")
PART = None  # document part, set in main(), for hyperlink relationships


def rpr(bold=False, italic=False, size=None, color=None, font=None):
    r = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        f.set(W(a), font or FONT)
    r.append(f)
    if bold:
        for t in ("w:b", "w:bCs"):
            r.append(OxmlElement(t))
    if italic:
        for t in ("w:i", "w:iCs"):
            r.append(OxmlElement(t))
    if color:
        c = OxmlElement("w:color"); c.set(W("w:val"), color); r.append(c)
    if size:
        for t in ("w:sz", "w:szCs"):
            s = OxmlElement(t); s.set(W("w:val"), str(int(size * 2))); r.append(s)
    return r


def add_runs(p, text, size=None, bold_all=False, italic_all=False):
    text = N.render(text)
    assert "{{" not in text, text
    for tok in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", text):
        if not tok:
            continue
        b, i = bold_all, italic_all
        if tok.startswith("**"):
            tok, b = tok[2:-2], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            tok, i = tok[1:-1], True
        for chunk in URL_RE.split(tok):
            if not chunk:
                continue
            target = p
            if URL_RE.fullmatch(chunk) and PART is not None:
                rid = PART.relate_to(chunk, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                     is_external=True)
                target = OxmlElement("w:hyperlink"); target.set(qn("r:id"), rid); p.append(target)
            for piece in re.split(f"([{FALLBACK_CHARS}])", chunk):
                if not piece:
                    continue
                r = OxmlElement("w:r")
                r.append(rpr(b, i, size, font=MINUS_FONT if piece in FALLBACK_CHARS else None))
                t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = piece
                r.append(t)
                target.append(r)


def set_spacing(p, before=None, after=None, line=None):
    ppr = p.find(W("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr"); p.insert(0, ppr)
    sp = ppr.find(W("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        # schema order: pStyle, keepNext, keepLines, ..., numPr, ..., spacing, ind, jc, rPr
        anchor = None
        for tag in ("w:ind", "w:jc", "w:rPr"):
            anchor = ppr.find(W(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(sp)
        else:
            ppr.append(sp)
    if before is not None:
        sp.set(W("w:before"), str(before))
    if after is not None:
        sp.set(W("w:after"), str(after))
    if line is not None:
        sp.set(W("w:line"), str(line)); sp.set(W("w:lineRule"), "auto")


def set_jc(p, val):
    ppr = p.find(W("w:pPr"))
    jc = ppr.find(W("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        rp = ppr.find(W("w:rPr"))
        if rp is not None:
            rp.addprevious(jc)
        else:
            ppr.append(jc)
    jc.set(W("w:val"), val)


def keep_next(p):
    ppr = p.find(W("w:pPr"))
    if ppr.find(W("w:keepNext")) is None:
        k = OxmlElement("w:keepNext")
        st = ppr.find(W("w:pStyle"))
        (st.addnext(k) if st is not None else ppr.insert(0, k))


def set_num(p, num_id):
    ppr = p.find(W("w:pPr"))
    np_ = ppr.find(W("w:numPr"))
    np_.find(W("w:numId")).set(W("w:val"), str(num_id))


def main():
    global PART
    d = docx.Document(str(TEMPLATE))
    PART = d.part
    body = d.element.body
    cp = d.core_properties
    cp.title, cp.author = content.TITLE, content.AUTHORS
    cp.subject = "Apart Research AI Collusion Research Sprint, October 2026"

    # --- prototypes from the template ----------------------------------------------------------
    P_H2 = copy.deepcopy(find_par(body, "1. Introduction"))
    P_H3 = copy.deepcopy(find_par(body, "Limitations"))
    P_BODY = copy.deepcopy(find_par(body, "What problem are you addressing"))
    P_LIST = copy.deepcopy(find_par(body, "[First contribution"))
    P_REF = copy.deepcopy(find_par(body, "[Reference 1]"))

    # --- title block ---------------------------------------------------------------------------
    header = d.tables[0]
    title_p = find_par(header._tbl, "PROJECT TITLE")
    title_run = [r for r in title_p.iter(W("w:r")) if text_of(r) == "PROJECT TITLE"][0]
    title_run.find(W("w:t")).text = content.TITLE
    for r in title_run.iter(W("w:sz")):
        r.set(W("w:val"), "32")
    for r in title_run.iter(W("w:szCs")):
        r.set(W("w:val"), "32")
    for r in title_p.iter(W("w:r")):  # footnote mark at the same size
        for s in r.iter(W("w:sz")):
            s.set(W("w:val"), "32")

    # authors: replace the 2x3 placeholder table with one centred paragraph
    author_tbl = [t for t in header._tbl.iter(W("w:tbl")) if "Author name 1" in text_of(t)][-1]
    ap = copy.deepcopy(P_BODY)
    strip_runs(ap)
    set_jc(ap, "center"); set_spacing(ap, before=60, after=120, line=240)
    add_runs(ap, content.AUTHORS, bold_all=True)
    br = OxmlElement("w:r"); br.append(OxmlElement("w:br")); ap.append(br)
    add_runs(ap, content.AFFIL)
    author_tbl.addprevious(ap)
    author_tbl.getparent().remove(author_tbl)

    # abstract
    abs_p = find_par(header._tbl, "Summarize your project")
    strip_runs(abs_p)
    rp = abs_p.find(W("w:pPr")).find(W("w:rPr"))
    if rp is not None and rp.find(W("w:u")) is not None:
        rp.remove(rp.find(W("w:u")))
    add_runs(abs_p, content.ABSTRACT, size=10)
    set_spacing(abs_p, after=60, line=240)

    # --- remove the guidance box and all template body content ---------------------------------
    d.tables[1]._tbl.getparent().remove(d.tables[1]._tbl)
    start = find_par(body, "1. Introduction")
    sect = body.find(W("w:sectPr"))
    el = start
    while el is not None and el is not sect:
        nxt = el.getnext()
        body.remove(el)
        el = nxt

    def emit(el):
        sect.addprevious(el)
        return el

    def par(proto, text, **kw):
        p = copy.deepcopy(proto)
        strip_runs(p)
        add_runs(p, text, **kw)
        return emit(p)

    list_counter = {"numbered_used": False}
    for blk in content.B:
        kind = blk[0]
        if kind == "h2":
            p = par(P_H2, blk[1]); set_spacing(p, before=180, after=40, line=260)
        elif kind == "h3":
            p = par(P_H3, blk[1], size=12); set_spacing(p, before=100, after=30, line=260)
        elif kind == "p":
            p = par(P_BODY, blk[1], size=BODY_PT); set_jc(p, "left" if "https://" in blk[1] else "both")
            set_spacing(p, after=60, line=245)
        elif kind in ("bullets", "numbered"):
            num = 3 if kind == "bullets" else (5 if not list_counter["numbered_used"] else 1)
            if kind == "numbered":
                list_counter["numbered_used"] = True
            for item in blk[1]:
                p = par(P_LIST, item, size=BODY_PT); set_num(p, num)
                set_jc(p, "left" if re.search(r"\S{35,}", N.render(item)) else "both")  # long paths justify badly
                set_spacing(p, after=40, line=252)
        elif kind == "fig":
            _, fname, width, caption = blk
            p = copy.deepcopy(P_BODY); strip_runs(p); emit(p)
            set_jc(p, "center"); set_spacing(p, before=80, after=40, line=240); keep_next(p)
            from docx.text.paragraph import Paragraph
            Paragraph(p, d._body).add_run().add_picture(str(FIGS / fname), width=Inches(width))
            c = par(P_BODY, caption, size=9.5); set_jc(c, "both"); set_spacing(c, after=100, line=230)
        elif kind == "table":
            _, caption, head, rows, widths = blk
            c = par(P_BODY, caption, size=9.5); set_jc(c, "both"); set_spacing(c, before=80, after=60, line=240)
            keep_next(c)
            t = d.add_table(rows=1 + len(rows), cols=len(head))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl = t._tbl
            tblPr = tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge, sz in (("top", 12), ("bottom", 12), ("insideH", 4)):
                e = OxmlElement(f"w:{edge}"); e.set(W("w:val"), "single"); e.set(W("w:sz"), str(sz))
                e.set(W("w:space"), "0"); e.set(W("w:color"), "000000"); borders.append(e)
            tblPr.find(W("w:tblLook")).addprevious(borders)
            layout = OxmlElement("w:tblLayout"); layout.set(W("w:type"), "fixed"); tblPr.find(W("w:tblLook")).addprevious(layout)
            grid = tbl.find(W("w:tblGrid"))
            for j, gc in enumerate(grid.findall(W("w:gridCol"))):
                gc.set(W("w:w"), str(int(widths[j] * 1440)))
            for j, w in enumerate(widths):
                for row in t.rows:
                    row.cells[j].width = Inches(w)
            for i, cells in enumerate([head] + rows):
                for j, txt in enumerate(cells):
                    cp = t.rows[i].cells[j].paragraphs[0]._p
                    add_runs(cp, txt, size=9, bold_all=(i == 0))
                    ppr = cp.get_or_add_pPr()
                    set_spacing(cp, before=20, after=20, line=240)

            # no row splits across pages, header repeats, and the table stays in one piece
            trs = tbl.findall(W("w:tr"))
            for i, tr in enumerate(trs):
                trPr = OxmlElement("w:trPr"); trPr.append(OxmlElement("w:cantSplit"))
                if i == 0:
                    trPr.append(OxmlElement("w:tblHeader"))
                tr.insert(0, trPr)
                if i < len(trs) - 1:
                    for cp_ in tr.iter(W("w:p")):
                        keep_next(cp_)
            sect.addprevious(tbl)
            sp = copy.deepcopy(P_BODY); strip_runs(sp); set_spacing(sp, after=80, line=200); emit(sp)
        elif kind == "refs":
            for ref in blk[1]:
                p = copy.deepcopy(P_BODY); strip_runs(p)
                ppr = p.find(W("w:pPr"))
                ind = OxmlElement("w:ind"); ind.set(W("w:left"), "360"); ind.set(W("w:hanging"), "360")
                ppr.find(W("w:rPr")).addprevious(ind)
                add_runs(p, ref, size=10); set_spacing(p, after=60, line=240); emit(p)
        elif kind == "pagebreak":
            p = copy.deepcopy(P_BODY); strip_runs(p)
            r = OxmlElement("w:r"); b = OxmlElement("w:br"); b.set(W("w:type"), "page"); r.append(b); p.append(r)
            emit(p)
        else:
            raise ValueError(kind)

    OUT.mkdir(exist_ok=True)
    out = OUT / "undercover_agents.docx"
    d.save(str(out))

    # provenance list for reviewers
    lines = ["# Number sources", "",
             "Every placeholder resolved in the paper, with the value printed. `ci:`/`v:`/`pp:`/`ppci:` keys are",
             "claims in `sting_results/claims.json`; `x:` keys are defined in `paper/paper_numbers.py` (source",
             "file noted there).", "", "| placeholder | printed |", "|---|---|"]
    lines += [f"| `{k}` | {v} |" for k, v in sorted(N.USED.items())]
    (OUT / "number_sources.md").write_text("\n".join(lines) + "\n")

    soffice = os.environ.get("SOFFICE", "soffice")
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(out)],
                       check=True, capture_output=True, timeout=180)
    except (FileNotFoundError, subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
        print("PDF not built:", e)
    print(out, f"{len(N.USED)} numbers resolved")
    abstract_words = len(N.render(content.ABSTRACT).split())
    print("abstract words:", abstract_words)


if __name__ == "__main__":
    main()
